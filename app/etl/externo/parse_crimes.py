"""Leitura da planilha de crimes por UF (aba consolidada "dados completos").

A aba é larga (um indicador por coluna); o parser devolve formato longo, um registro por
UF × ano × indicador, preservando o valor bruto da célula e o motivo de um valor ausente
("..." e célula vazia não são zero).
"""

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import BinaryIO

import openpyxl

from app.enrichment.externo.taxonomy import UF_BY_NAME, fold, parse_br_int

SHEET = "dados completos"

# cabeçalho (sem acento, minúsculo) -> indicador canônico
INDICADORES = {
    "estupro": "estupro",
    "delitos envolvendo drogas (trafico)": "trafico_drogas",
    "delitos envolvendo drogas (posse e uso)": "posse_drogas",
    "porte ilegal de arma de fogo": "porte_arma",
    "roubos(total)": "roubo_total",
    "homicidio doloso": "homicidio_doloso",
    "latrocinios": "latrocinio",
    "lesao corporal seguida de morte": "lesao_corporal_morte",
    "tentativa de homicidio": "tentativa_homicidio",
    "suicidio": "suicidio",
    "mortes a esclarecer": "mortes_a_esclarecer",
    "roubo a instituicao financeira": "roubo_instituicao_financeira",
    "roubo de carga": "roubo_carga",
    "roubo de veiculo": "roubo_veiculo",
    "roubo em estabelecimento comercial ou de servico": "roubo_estabelecimento_comercial",
    "populacao": "populacao",
}
_PRESOS = "presos no sistema penitenciario"
_YEAR_SHEET = re.compile(r"^\s*(\d{4})\.?\s*$")


@dataclass
class CrimeRecord:
    uf: str
    ano: int
    indicador: str
    valor: int | None
    situacao: str
    valor_bruto: str | None
    ressalva: str | None
    aba_origem: str
    linha_origem: int
    qa_flags: list[str] = field(default_factory=list)


@dataclass
class CrimeRejection:
    linha: int
    motivo: str
    conteudo: list[str]


@dataclass
class ParsedCrimes:
    registros: list[CrimeRecord] = field(default_factory=list)
    rejeicoes: list[CrimeRejection] = field(default_factory=list)
    colunas_ignoradas: list[str] = field(default_factory=list)


def _indicador(header: object) -> str | None:
    name = fold(header or "")
    if name.startswith(_PRESOS):
        return "presos_sistema_penitenciario"
    return INDICADORES.get(name)


def _year_notes(wb: openpyxl.Workbook) -> dict[int, str]:
    """Notas iniciadas por '*' no topo das abas anuais (ex.: dados parciais de 2020)."""
    notes: dict[int, str] = {}
    for ws in wb.worksheets:
        match = _YEAR_SHEET.match(ws.title)
        if not match:
            continue
        for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
            for value in row:
                if isinstance(value, str) and value.strip().startswith("*"):
                    notes[int(match.group(1))] = value.strip().lstrip("*").strip()
    return notes


def parse_crimes_workbook(source: Path | BinaryIO) -> ParsedCrimes:
    wb = openpyxl.load_workbook(source, data_only=True)
    if SHEET not in wb.sheetnames:
        raise ValueError(f"aba '{SHEET}' não encontrada ({wb.sheetnames})")
    ws = wb[SHEET]
    notes = _year_notes(wb)
    parsed = ParsedCrimes()

    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    estado_col, ano_col = header.index("Estado"), header.index("Ano")
    columns: dict[int, str] = {}
    for index, name in enumerate(header):
        if index in (estado_col, ano_col):
            continue
        indicador = _indicador(name)
        if indicador:
            columns[index] = indicador
        else:
            parsed.colunas_ignoradas.append(str(name).strip())

    for linha, row in enumerate(rows, start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        conteudo = ["" if v is None else str(v) for v in row]
        uf = UF_BY_NAME.get(fold(row[estado_col] or ""))
        if uf is None:
            parsed.rejeicoes.append(CrimeRejection(linha, "estado_desconhecido", conteudo))
            continue
        ano = row[ano_col]
        if not isinstance(ano, int | float) or not 1990 <= int(ano) <= 2100:
            parsed.rejeicoes.append(CrimeRejection(linha, "ano_invalido", conteudo))
            continue
        ano = int(ano)
        for index, indicador in columns.items():
            raw = row[index]
            valor, situacao = parse_br_int(raw)
            parsed.registros.append(
                CrimeRecord(
                    uf=uf,
                    ano=ano,
                    indicador=indicador,
                    valor=valor,
                    situacao=situacao,
                    valor_bruto=None if raw is None else str(raw),
                    ressalva=notes.get(ano),
                    aba_origem=SHEET,
                    linha_origem=linha,
                )
            )
    return parsed
